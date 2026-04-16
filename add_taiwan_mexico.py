import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ===== 既存ファイルを読み込む =====
path = "/home/user/congenial-journey/overseas_buyers_BR380JG3.xlsx"
wb = openpyxl.load_workbook(path)

HEADER_BG = "1F4E79"
HEADER_FG = "FFFFFF"
S3_BG     = "FF0000"
S2_BG     = "FF9900"
S1_BG     = "92D050"
ALT_BG    = "EBF3FB"
WHITE     = "FFFFFF"
TAIWAN_H  = "2E75B6"   # 台湾シートヘッダー：青
MEXICO_H  = "538135"   # メキシコシートヘッダー：緑

def hdr_fill(c): return PatternFill("solid", fgColor=c)
def cell_fill(c): return PatternFill("solid", fgColor=c)
thin = Side(style="thin", color="CCCCCC")
thin_border = Border(left=thin, right=thin, top=thin, bottom=thin)
priority_colors = {"★★★": S3_BG, "★★": S2_BG, "★": S1_BG}

def build_sheet(wb, title, hdr_color, headers, data, col_widths):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for col in range(1, len(headers)+1):
        c = ws.cell(row=1, column=col)
        c.fill = hdr_fill(hdr_color)
        c.font = Font(color=HEADER_FG, bold=True, size=10)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = thin_border
    ws.row_dimensions[1].height = 30

    for i, row in enumerate(data, 2):
        ws.append(list(row))
        bg = ALT_BG if i % 2 == 0 else WHITE
        prio = row[1] if len(row) > 1 else ""
        for col in range(1, len(headers)+1):
            c = ws.cell(row=i, column=col)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = thin_border
            if col == 2 and prio in priority_colors:
                c.fill = cell_fill(priority_colors[prio])
                c.font = Font(bold=True, color="FFFFFF")
                c.alignment = Alignment(horizontal="center", vertical="center")
            else:
                c.fill = cell_fill(bg)
        ws.row_dimensions[i].height = 52

    for col, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws

# ============================================================
# Sheet 4: 台湾 – BR380JGポテンシャルバイヤー
# ============================================================
tw_headers = [
    "No", "優先度", "会社名（中文）", "会社名（日本語）", "所在地（台湾）",
    "ウェブサイト", "メールアドレス", "電話番号", "WhatsApp",
    "Facebook / SNS", "日本製機械取扱", "BR380JG関連", "備考"
]
tw_data = [
    (1, "★★★",
     "鑫輝機械有限公司",
     "シンフェイ機械",
     "基隆市七堵区",
     "https://www.bmcl.com.tw",
     "bmcl66@mail2000.com.tw",
     "02-24579141", "",
     "",
     "◎ 日本TOKU製新品破砕機・中古建機を扱う",
     "日本TOKU製新品クラッシャー（破碎機）販売実績あり",
     "日本TOKUの新品クラッシャー代理店かつ中古ショベル・ブルドーザー等も扱う。BR380JG（移動式ジョークラッシャー）の需要が高い。最優先ターゲット。"),

    (2, "★★★",
     "互益機械有限公司（互助重機）",
     "フーイー機械",
     "高雄市大社区",
     "https://www.huyih.com",
     "huyih.hy@msa.hinet.net",
     "07-3535686", "",
     "http://huyih.blogspot.com/",
     "◎ 外匯（日本）中古重機専門",
     "日本からの中古ショベル輸入実績あり",
     "「外匯中古重機」専門業者（外匯＝海外からの輸入）。日本製コマツ・日立・CAT等の中古ショベル・ブルドーザーを買取・販売。破砕機ニーズも高い。"),

    (3, "★★",
     "台松堆高機股份有限公司",
     "台松フォークリフト（Komatsu台湾代理店）",
     "台北市",
     "http://www.taiwan-komatsu.tw",
     "Webフォームより", "", "",
     "https://www.facebook.com/taiwan.komatsu/",
     "◎ コマツ台湾正規代理店",
     "コマツ全製品取扱（フォークリフト中心）",
     "日本小松台湾区総代理店。主にフォークリフト中心だが建設機械・コマツ全製品を取扱。BR380JG需要の把握・紹介元として重要。"),

    (4, "★★",
     "鴻聖重機有限公司",
     "ホンセン重機",
     "彰化県",
     "https://www.hs-digger.com.tw",
     "Webフォームより", "", "",
     "",
     "◎ 中古ショベル・建機買取販売",
     "中古建機輸入業者",
     "1997年創業。工地（建設現場）・外匯（日本）からの中古ショベル・ブルドーザー等を専門に売買。日本からの輸入実績あり。"),

    (5, "★★",
     "洪鐵重機械股份有限公司",
     "ホンティエ重機",
     "台湾（北部）",
     "https://www.hung-tieh.com.tw",
     "Webフォームより", "", "",
     "",
     "◎ 日本製中古重機",
     "日本製中古建機輸入・販売",
     "日本製中古ショベル・ローダー・ロードローラー・発電機・堆高機（フォークリフト）等を扱う。日本からの輸入メイン。"),

    (6, "★★",
     "鉅工堆高機",
     "キョコウフォークリフト",
     "台湾",
     "https://www.the-giant.com.tw",
     "Webフォームより", "", "",
     "",
     "◎ コマツ代理権取得",
     "コマツ電動フォークリフト代理",
     "コマツ電動式フォークリフト代理権取得業者。コマツネットワーク活用でBR380JG導入実績業者への紹介が期待できる。"),

    (7, "★★",
     "台明重工",
     "タイメン重工",
     "台湾",
     "http://www.backhoe.com.tw",
     "Webフォームより", "", "",
     "",
     "◎ 中古ショベル売買",
     "中古建機売買",
     "中古ショベル（怪手・挖土機）専門。日本輸入品を扱う可能性あり。"),

    (8, "★★",
     "勵國重機有限公司",
     "リーグオ重機",
     "新北市",
     "Webサイト確認中",
     "Webフォームより", "", "",
     "https://www.facebook.com/UsedShovelLoader/",
     "◎ コマツ中古重機のFacebook投稿あり",
     "Komatsu WA450-3等の中古重機をFacebook上で販売",
     "Facebook上でKomatsu WA450-3等の中古重機を販売。コマツファン層。BR380JGに興味を持つ可能性が高い。"),

    (9, "★",
     "Hitachi Construction Machinery Taiwan",
     "日立建機台湾",
     "台湾（公式代理店）",
     "https://www.hitachicm.com.tw",
     "Webフォームより", "", "",
     "",
     "◎ 日立建機台湾公式",
     "日立製品のみ",
     "日立建機の台湾公式代理店。競合だが、建設・採石業界へのネットワークを通じてBR380JGのニーズを探る情報源として活用可。"),

    (10, "★",
     "豐悅企業有限公司",
     "ホンユエ企業",
     "台湾",
     "https://www.fycmg.com",
     "Webフォームより", "", "",
     "",
     "◎ 台湾柳工代理店・重機全般",
     "建設機械全般",
     "台湾柳工（LiuGong）代理店。ショベル・ローダー・プッシャー・フォークリフト・圧路機・ブルドーザー等取扱。破砕機ニーズ調査先として有望。"),
]

tw_widths = [4, 8, 22, 20, 14, 35, 28, 16, 12, 42, 22, 20, 55]
build_sheet(wb, "台湾バイヤーリスト", TAIWAN_H, tw_headers, tw_data, tw_widths)

# ============================================================
# Sheet 5: メキシコ – 日本製中古建機（CAT中心）バイヤー
# ============================================================
mx_headers = [
    "No", "優先度", "会社名", "所在地（メキシコ）",
    "ウェブサイト", "メールアドレス", "電話番号", "WhatsApp",
    "Facebook / SNS", "日本製機械取扱", "CATメイン", "備考"
]
mx_data = [
    (1, "★★★",
     "Maquinas Diesel S.A. de C.V. (MADISA-CAT)",
     "Av. Industriales del Poniente 2300, Santa Catarina, Nuevo León, 66350",
     "https://madisa.com.mx",
     "Webフォームより",
     "+52 81 8400 2000", "",
     "",
     "◎ CAT公認代理店",
     "YES – CAT専門・メキシコ最大",
     "メキシコ最大のCaterpillar正規代理店（1946年創業）。建設・鉱山・石油ガス・農業向け。Telsmith社製クラッシャーも扱う。日本製CAT中古機の輸入購買実績が期待できる最重要ターゲット。"),

    (2, "★★★",
     "Ritchie Bros. Auctioneers de Mexico",
     "Polotitlan, Estado de Mexico（競売場）",
     "https://www.rbauction.com",
     "Webフォームより",
     "Webより確認", "",
     "",
     "◎ 国際オークション・輸入実績多数",
     "YES – CAT輸入84件",
     "メキシコ向けCAT掘削機を84件輸入した実績あり（輸入統計）。メキシコ国内で定期競売を開催。日本製中古CAT建機を競売出品すれば直接リーチ可能。"),

    (3, "★★★",
     "Maquinaria Wiebe KM 24 S.A. de C.V.",
     "Km 24 Carretera Cuauhtémoc a Álvaro Obregón, Cuauhtémoc, Chihuahua",
     "https://www.maquinariaw.com",
     "awiebe71@hotmail.com\nventas@maquinariaw.com.mx",
     "+52 625 134 4523", "",
     "https://www.facebook.com/maquinariaw/",
     "◎ 中古CAT・建機売買専門",
     "YES – CAT 320C・CAT 349FL等多数",
     "メキシコ最大級の中古建機ディーラー（MachineryTraderに1,500件超の出品）。CAT掘削機・クレーン・フォークリフト等を扱う。米国・日本からの輸入品を販売。メール・Facebookで連絡可。"),

    (4, "★★",
     "Tracsa S.A.P.I. de C.V. (Grupo Tracsa)",
     "Periferico Sur 7800, Santa Maria Tequepexpan, 45600 Tlaquepaque, Jalisco",
     "https://tracsa.com.mx",
     "Webフォームより",
     "Webより確認", "",
     "https://www.linkedin.com/company/tracsacat",
     "◎ CAT公認代理店（西部・バヒオ地方）",
     "YES – CAT専門",
     "メキシコ西部・バヒオ地方のCAT正規代理店（1974年〜）。14拠点。年商5億ドル超（2025年）。建設・農業・鉱業・物流。新旧CAT機械取扱。LinkedIn: Tracsa Cat。"),

    (5, "★★",
     "TMR – Tractores y Maquinaria Real",
     "Carretera Miguel Alemán 102 (Km 12.5), San Nicolás de los Garza, Nuevo León",
     "https://www.tmr.com.mx",
     "Webフォームより",
     "+52 81 8327 0948", "",
     "https://www.linkedin.com/in/tractores-y-maquinaria-real/",
     "◎ CAT公認代理店（北部）",
     "YES – CAT専門",
     "モンテレイ・サルティーヨ・トレオン拠点のCAT公認ディーラー。新旧重機・建機の販売・レンタル。Dynapac代理店も兼任。北部メキシコ最大拠点の一つ。"),

    (6, "★★",
     "Empresas MATCO S.A. de C.V. (Matco Cat)",
     "Sufragio Efectivo Norte 870, Col. Centro, Ciudad Obregón, Sonora\n他: Mexicali / Mazatlán / Culiacán / Los Mochis / Hermosillo",
     "https://matco.com.mx",
     "m***@matco.com.mx (ZoomInfo経由要確認)",
     "Webより確認", "",
     "https://www.linkedin.com/company/empresas-matco-s-a-de-c-v-/",
     "◎ CAT公認代理店（西部・北西部）",
     "YES – CAT 300モデル以上",
     "ソノラ州・バハカリフォルニア州等のCAT正規代理店。300モデル以上の機械カタログ。農業・建設・鉱業・海事。Astec Telsmith製クラッシャーも販売。"),

    (7, "★★",
     "Heavy Machine México",
     "メキシコ（詳細要確認）",
     "https://heavymachine.com.mx",
     "Webフォームより",
     "Webより確認", "",
     "",
     "◎ 中古CAT・建機全般",
     "YES – CAT含む複数ブランド",
     "CAT・Case・Bobcat・JD・JCB等の中古重機を扱うメキシコの建機ディーラー。日本製中古機も取扱可能性あり。"),

    (8, "★",
     "MarketBook México",
     "メキシコ全土（オンラインマーケット）",
     "https://www.marketbook.mx",
     "Webフォームより",
     "", "",
     "",
     "◎ 中古建機マーケットプレイス",
     "YES – 各ブランド掲載",
     "メキシコ国内最大級の中古農業機械・建機オンラインマーケット。出品掲載でメキシコのバイヤーに直接リーチ可能。"),
]

mx_widths = [4, 8, 30, 30, 30, 30, 18, 12, 38, 22, 18, 60]
build_sheet(wb, "メキシコバイヤーリスト", MEXICO_H, mx_headers, mx_data, mx_widths)

wb.save(path)
print(f"Updated: {path}")
